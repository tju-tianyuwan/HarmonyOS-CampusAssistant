"""Bounded document decoding, run in a worker thread by the upload route."""
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile
import base64
from functools import lru_cache
from threading import Lock

from fastapi import HTTPException

from ..config import settings

ocr_lock = Lock()


@lru_cache(maxsize=1)
def local_ocr():
    from rapidocr_onnxruntime import RapidOCR
    return RapidOCR(intra_op_num_threads=2, inter_op_num_threads=1)


def recognize_image(data: bytes, mime: str = "image/png") -> str:
    if settings.document_ocr_provider == "local":
        with ocr_lock:
            result, _ = local_ocr()(data)
        return "\n".join(row[1] for row in result or [])
    if not settings.llm_vision_model or not settings.llm_api_key:
        raise HTTPException(422, "图片或扫描页需要配置 LLM_VISION_MODEL 和 LLM_API_KEY")
    from openai import OpenAI
    try:
        with OpenAI(base_url=settings.llm_base_url, api_key=settings.llm_api_key,
                    timeout=45, max_retries=0) as client:
            response = client.chat.completions.create(model=settings.llm_vision_model, messages=[{
                "role": "user", "content": [
                    {"type": "text", "text": "提取图片中的文字、公式和表格为 Markdown。只转录可见内容，不解释、不补充。空白图片返回空字符串。"},
                    {"type": "image_url", "image_url": {"url": f"data:{mime};base64," + base64.b64encode(data).decode("ascii")}},
                ]}])
        return response.choices[0].message.content or ""
    except Exception as exc:
        raise HTTPException(502, "图片文字识别失败，未保存不完整的文档，请检查视觉模型配置后重试") from exc


def bounded_image(data: bytes) -> bytes:
    import fitz
    with fitz.open(stream=data) as image:
        page = image[0]
        scale = min(2, 1800 / max(page.rect.width, page.rect.height))
        return page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False).tobytes("png")


def extract(filename: str, data: bytes, progress=None) -> str:
    suffix = Path(filename).suffix.lower()
    try:
        if suffix in (".md", ".txt"):
            text = data.decode("utf-8-sig")
        elif suffix == ".pdf":
            import fitz
            with fitz.open(stream=data, filetype="pdf") as document:
                if document.page_count > settings.document_max_pages:
                    raise HTTPException(413, "PDF 页数超出限制")
                pages = []
                for index, page in enumerate(document):
                    content = page.get_text(sort=True)
                    if not content.strip():
                        scale = min(2, 1800 / max(page.rect.width, page.rect.height))
                        content = recognize_image(page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False).tobytes("png"))
                    else:
                        tables = page.find_tables()
                        content += "\n" + "\n".join(table.to_markdown() for table in tables.tables)
                        for image in page.get_image_info(xrefs=True):
                            if image.get("xref", 0) > 0 and image["width"] >= 120 and image["height"] >= 80:
                                content += "\n" + recognize_image(bounded_image(document.extract_image(image["xref"])["image"]))
                    pages.append(content)
                    if progress: progress(index + 1, document.page_count)
                if not any(page.strip() for page in pages):
                    raise HTTPException(422, "文档没有可识别的内容")
                text = "\n\n".join(f"## 第 {i + 1} 页\n{page}" for i, page in enumerate(pages))
        elif suffix == ".pptx":
            from pptx import Presentation
            with ZipFile(BytesIO(data)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 100 * 1024 * 1024:
                    raise HTTPException(413, "PPTX 解压大小超出限制")
            slides = Presentation(BytesIO(data)).slides
            if len(slides) > settings.document_max_pages:
                raise HTTPException(413, "幻灯片页数超出限制")
            pages = []
            def shapes_text(shapes):
                parts = []
                for shape in shapes:
                    if shape.shape_type == 6:
                        parts.extend(shapes_text(shape.shapes))
                    if shape.has_text_frame:
                        parts.append(shape.text)
                    if shape.has_table:
                        rows = [[cell.text.replace("|", "\\|").replace("\n", "<br>") for cell in row.cells] for row in shape.table.rows]
                        if rows:
                            parts.append("\n".join([" | ".join(rows[0]), " | ".join(["---"] * len(rows[0]))] + [" | ".join(row) for row in rows[1:]]))
                    if shape.has_chart:
                        chart = shape.chart
                        if chart.has_title: parts.append(chart.chart_title.text_frame.text)
                        for plot in chart.plots:
                            labels = [str(category.label) for category in plot.categories] if hasattr(plot, "categories") else []
                            for series in plot.series:
                                parts.append(str(series.name))
                                parts.extend(f"{labels[i] if i < len(labels) else i + 1}: {value}" for i, value in enumerate(series.values))
                    if shape.shape_type == 13:
                        parts.append(recognize_image(bounded_image(shape.image.blob)))
                return parts
            for index, slide in enumerate(slides):
                parts = []
                parts.extend(shapes_text(slide.shapes))
                if slide.has_notes_slide:
                    parts.append(slide.notes_slide.notes_text_frame.text)
                pages.append("\n".join(parts))
                if progress: progress(index + 1, len(slides))
            if not any(page.strip() for page in pages):
                raise HTTPException(422, "幻灯片没有可识别的内容")
            text = "\n\n".join(f"## 第 {i + 1} 页\n{page}" for i, page in enumerate(pages))
        else:
            raise HTTPException(415, "支持 PDF、PPTX、Markdown 和 TXT")
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(422, "文件损坏或编码不支持，请使用 UTF-8 文本或有效文档") from exc
    if not text.strip() or len(text) > 200000:
        raise HTTPException(422, "提取内容为空或超过 20 万字符")
    return text
