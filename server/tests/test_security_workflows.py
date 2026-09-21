import asyncio
import json
import unittest
from datetime import datetime, timedelta
from unittest.mock import AsyncMock, patch

from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient
from sqlalchemy.pool import StaticPool
from sqlmodel import SQLModel, Session, create_engine, select

from app.db import get_db
from app.models import (Account, AccountCode, AuthSession, ClassCourse, CourseSession, EmailCode, Membership,
                        MeetingMember, MeetingMessage, MeetingRoom, Note, Outline, PersonalWorkspace,
                        SessionResource, TranscriptSegment, User)
from app.routers import auth, courses, sessions, notes, chat, analytics, practice, knowledge, meetings
from app.services import rag
from app.services import meetings as scheduler
from app.services.security import hash_password, token_hash


class WorkflowsTest(unittest.TestCase):
    def setUp(self):
        self.engine = create_engine('sqlite://', connect_args={'check_same_thread': False}, poolclass=StaticPool)
        SQLModel.metadata.create_all(self.engine)
        app = FastAPI()
        for router in (auth, courses, sessions, notes, chat, analytics, practice, knowledge, meetings):
            app.include_router(router.router, prefix='/api/v1')
        def database():
            with Session(self.engine) as db:
                yield db
        app.dependency_overrides[get_db] = database
        self.client = TestClient(app)
        self.embed_patch = patch.object(rag.settings, 'embed_api_key', '')
        self.embed_patch.start()
        with Session(self.engine) as db:
            db.add_all([User(id=1, name='Teacher', role='teacher'), User(id=2, name='Student'), User(id=3, name='Outsider')])
            db.add(ClassCourse(id=1, name='Data structures', class_name='A', teacher_id=1, invite_code='123456'))
            db.add_all([Membership(user_id=1, class_course_id=1), Membership(user_id=2, class_course_id=1)])
            db.add_all([AuthSession(token_hash=token_hash(f'token-{i}'), user_id=i,
                                    expires_at=datetime.utcnow() + timedelta(hours=1)) for i in range(1, 4)])
            db.add(CourseSession(id=1, class_course_id=1, title='Trees', creator_id=1))
            db.commit()
        self.client.headers['Authorization'] = 'Bearer token-1'

    def tearDown(self):
        self.client.close()
        self.embed_patch.stop()
        self.engine.dispose()

    def as_user(self, user):
        self.client.headers['Authorization'] = f'Bearer token-{user}'

    def test_teacher_materials_are_course_scoped_and_exclude_personal_notes(self):
        with Session(self.engine) as db:
            db.add(Outline(session_id=1, class_course_id=1, owner_id=-1, markdown='Official', status='pending'))
            db.add(Outline(session_id=1, class_course_id=1, owner_id=2, markdown='Private'))
            db.add(SessionResource(session_id=1, class_course_id=1, uploader_id=1, filename='lesson.txt', content='Lesson'))
            db.add(ClassCourse(id=2, name='Other', class_name='B', teacher_id=1, invite_code='654321'))
            db.add(CourseSession(id=2, class_course_id=2, title='Other lesson', creator_id=1))
            db.add(Outline(session_id=2, class_course_id=2, owner_id=-1, markdown='Other'))
            db.add(SessionResource(session_id=2, class_course_id=2, uploader_id=1, filename='other.txt', content='Other'))
            db.commit()
        response = self.client.get('/api/v1/sessions/materials?class_course_id=1')
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.headers['cache-control'], 'no-store')
        self.assertEqual(len(response.json()), 2)
        self.assertEqual({item['kind'] for item in response.json()}, {'outline', 'resource'})
        self.assertTrue(all(item['session_id'] == 1 for item in response.json()))
        for user in (2, 3):
            self.as_user(user)
            self.assertEqual(self.client.get('/api/v1/sessions/materials?class_course_id=1').status_code, 403)
        self.client.headers.clear()
        self.assertEqual(self.client.get('/api/v1/sessions/materials?class_course_id=1').status_code, 401)

    def test_resource_edit_requires_manager_and_matching_session(self):
        with Session(self.engine) as db:
            db.add(Outline(session_id=1, class_course_id=1, owner_id=-1, markdown='Official', status='published'))
            db.add(SessionResource(id=1, session_id=1, class_course_id=1, uploader_id=1, filename='lesson.txt', content='Original'))
            db.add(CourseSession(id=2, class_course_id=1, title='Next lesson', creator_id=1))
            db.commit()
        path = '/api/v1/sessions/1/resources/1'
        with patch.object(rag, 'add_document') as index:
            result = self.client.put(path, json={'content': 'Corrected lesson'})
            self.assertEqual(result.status_code, 200, result.text)
            self.assertEqual(result.json()['content'], 'Corrected lesson')
            index.assert_called_once()
        self.assertEqual(self.client.put('/api/v1/sessions/2/resources/1', json={'content': 'Wrong session'}).status_code, 404)
        self.assertEqual(self.client.put(path, json={'content': '  '}).status_code, 422)
        self.as_user(2)
        self.assertEqual(self.client.put(path, json={'content': 'Student edit'}).status_code, 403)
        self.assertEqual(self.client.get(path).json()['content'], 'Corrected lesson')

    def test_unauthenticated_and_impersonation(self):
        self.client.headers.clear()
        for path in ('/auth/users', '/courses?user_id=1', '/sessions?class_course_id=1', '/knowledge/1'):
            self.assertEqual(self.client.get('/api/v1' + path).status_code, 401, path)
        self.as_user(2)
        self.assertEqual(self.client.get('/api/v1/courses?user_id=1').status_code, 403)
        self.assertEqual(self.client.get('/api/v1/analytics/classroom?class_course_id=1&user_id=1').status_code, 403)

    def test_register_school_login_logout(self):
        body = dict(account_type='school', school_name='University', student_number='2026001',
                    name='New Teacher', password='strong-password', role='teacher')
        response = self.client.post('/api/v1/auth/register', json=body)
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()['user']['role'], 'teacher')
        self.client.headers['Authorization'] = 'Bearer ' + response.json()['access_token']
        self.assertEqual(self.client.get('/api/v1/auth/me').status_code, 200)
        self.assertEqual(self.client.post('/api/v1/auth/logout').status_code, 200)
        self.assertEqual(self.client.get('/api/v1/auth/me').status_code, 401)
        self.assertEqual(self.client.post('/api/v1/auth/login', json=body).status_code, 200)
        body['password'] = 'wrong-password'
        self.assertEqual(self.client.post('/api/v1/auth/login', json=body).status_code, 401)

    def test_registration_role_defaults_to_student_and_rejects_unknown_role(self):
        body = dict(account_type='school', school_name='University', student_number='2026002',
                    name='New Student', password='strong-password')
        response = self.client.post('/api/v1/auth/register', json=body)
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()['user']['role'], 'student')

        body['student_number'] = '2026003'
        body['role'] = 'admin'
        self.assertEqual(self.client.post('/api/v1/auth/register', json=body).status_code, 422)

    def test_personal_registration_consumes_code_and_provisions_private_course(self):
        with Session(self.engine) as db:
            db.add(EmailCode(email='student@example.com', code_hash=hash_password('123456'),
                             expires_at=datetime.utcnow() + timedelta(minutes=10)))
            db.commit()
        body = dict(name='Personal Teacher', email='student@example.com', password='strong-password',
                    code='123456', role='teacher')
        response = self.client.post('/api/v1/auth/register', json=body)
        self.assertEqual(response.status_code, 200, response.text)
        user = response.json()['user']
        self.assertEqual(user['role'], 'teacher')
        self.client.headers['Authorization'] = 'Bearer ' + response.json()['access_token']
        courses_response = self.client.get(f"/api/v1/courses?user_id={user['id']}").json()
        self.assertTrue(courses_response[0]['is_personal'])
        self.assertEqual(courses_response[0]['invite_code'], '')
        with Session(self.engine) as db:
            self.assertIsNone(db.get(EmailCode, 'student@example.com'))

    def test_audio_limits_idempotence_and_finish(self):
        path = '/api/v1/sessions/1/chunks/'
        self.as_user(2)
        self.assertEqual(self.client.post(path + '0', files={'file': ('a.pcm', b'1234')}).status_code, 403)
        self.as_user(1)
        self.assertEqual(self.client.post(path + '-1', files={'file': ('a.pcm', b'1234')}).status_code, 422)
        self.assertEqual(self.client.post(path + '0', files={'file': ('a.pcm', b'')}).status_code, 400)
        self.assertEqual(self.client.post(path + '0', files={'file': ('a.pcm', b'x' * (3 * 1024 * 1024 + 1))}).status_code, 413)
        from app.services.asr.base import ASRResult
        provider = AsyncMock()
        provider.transcribe_chunk.return_value = ASRResult(text='binary trees', start_ms=0, end_ms=4000)
        with patch.object(sessions, 'get_asr_provider', return_value=provider):
            first = self.client.post(path + '0', files={'file': ('a.pcm', b'1234')})
            self.assertEqual(first.status_code, 200, first.text)
            duplicate = self.client.post(path + '0', files={'file': ('a.pcm', b'5678')})
            self.assertEqual(first.json()['id'], duplicate.json()['id'])
            self.assertEqual(provider.transcribe_chunk.await_count, 1)
            self.client.post('/api/v1/sessions/1/finish')
            self.assertEqual(self.client.post(path + '1', files={'file': ('a.pcm', b'5678')}).status_code, 409)

    def test_personal_outline_does_not_replace_official(self):
        with Session(self.engine) as db:
            db.add(Outline(session_id=1, class_course_id=1, owner_id=-1, markdown='Official', status='published'))
            db.commit()
        self.as_user(2)
        response = self.client.put('/api/v1/sessions/1/outline?user_id=2', json={'markdown': 'My version'})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()['owner_id'], 2)
        self.assertEqual(self.client.post('/api/v1/sessions/1/outline/review', json={'user_id': 2}).status_code, 403)
        self.as_user(1)
        self.assertEqual(self.client.get('/api/v1/sessions/1/outline').json()['markdown'], 'Official')

    def test_outline_generation_conflict_keeps_teacher_edit(self):
        with Session(self.engine) as db:
            db.add(TranscriptSegment(session_id=1, seq=0, text='trees', start_ms=0, end_ms=4000))
            db.add(Outline(session_id=1, class_course_id=1, owner_id=-1, markdown='Original'))
            db.commit()
        async def edit_during_generation(*args):
            with Session(self.engine) as db:
                outline = db.exec(select(Outline)).first()
                outline.markdown = 'Teacher edit'
                db.commit()
            return 'Generated'
        with patch.object(sessions, 'chat', side_effect=edit_during_generation):
            self.assertEqual(self.client.post('/api/v1/sessions/1/outline/generate').status_code, 409)
        self.assertEqual(self.client.get('/api/v1/sessions/1/outline').json()['markdown'], 'Teacher edit')

    def test_course_outline_catalog_is_manager_only(self):
        with Session(self.engine) as db:
            db.add(Outline(session_id=1, class_course_id=1, owner_id=-1,
                           markdown='# AI course outline', status='generated'))
            db.add(Outline(session_id=1, class_course_id=1, owner_id=2,
                           markdown='# Personal outline', status='published'))
            db.commit()
        response = self.client.get('/api/v1/sessions/outlines?class_course_id=1')
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(len(response.json()), 1)
        self.assertEqual(response.json()[0]['session_title'], 'Trees')
        self.assertEqual(response.json()[0]['markdown'], '# AI course outline')
        with Session(self.engine) as db:
            db.add(PersonalWorkspace(user_id=2, class_course_id=1))
            db.commit()
        self.as_user(2)
        self.assertEqual(self.client.get('/api/v1/sessions/outlines?class_course_id=1').status_code, 403)

    def test_note_privacy_and_rag_withdrawal(self):
        self.as_user(2)
        note = self.client.post('/api/v1/notes', json=dict(class_course_id=1, owner_id=2, title='Trees', content='二叉树遍历')).json()
        with patch.object(notes, 'chat', new=AsyncMock(return_value=json.dumps(dict(relevance=8, correctness=8, structure=8)))):
            shared = self.client.post(f"/api/v1/notes/{note['id']}/share")
            self.assertEqual(shared.status_code, 200, shared.text)
        hits = self.client.post('/api/v1/knowledge/1/search', json={'question': '二叉树'}).json()['hits']
        self.assertTrue(hits)
        self.client.post(f"/api/v1/notes/{note['id']}/unshare")
        self.assertEqual(self.client.post('/api/v1/knowledge/1/search', json={'question': '二叉树'}).json()['hits'], [])
        self.as_user(1)
        self.assertEqual(self.client.put(f"/api/v1/notes/{note['id']}", json={'content': 'overwrite'}).status_code, 403)
        self.as_user(3)
        self.assertEqual(self.client.get('/api/v1/notes?class_course_id=1&user_id=3').status_code, 403)

    def test_upload_and_query_isolation(self):
        response = self.client.post('/api/v1/knowledge/1/upload', files={'file': ('trees.md', '# 二叉树\n中序遍历'.encode())})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertTrue(self.client.post('/api/v1/knowledge/1/search', json={'question': '二叉树'}).json()['hits'])
        self.as_user(3)
        self.assertEqual(self.client.post('/api/v1/knowledge/1/search', json={'question': '二叉树'}).status_code, 403)

    def test_lesson_resources_and_personal_notes_are_scoped(self):
        with Session(self.engine) as db:
            db.add(Outline(session_id=1, class_course_id=1, owner_id=-1,
                           markdown='# Trees', status='published'))
            db.commit()
        with patch.object(sessions, 'extract', return_value='# Trees\nCourse handout'):
            response = self.client.post('/api/v1/sessions/1/resources',
                                        files={'file': ('trees.md', b'course handout')})
        self.assertEqual(response.status_code, 200, response.text)
        resource_id = response.json()['id']

        self.as_user(2)
        note = self.client.post('/api/v1/notes', json={
            'class_course_id': 1, 'session_id': 1, 'owner_id': 2,
            'title': 'My tree notes', 'content': 'Traversal',
        })
        self.assertEqual(note.status_code, 200, note.text)
        lesson = self.client.get('/api/v1/sessions?class_course_id=1').json()[0]
        self.assertEqual(lesson['resource_count'], 1)
        self.assertEqual([row['title'] for row in lesson['personal_notes']], ['My tree notes'])
        self.assertEqual(self.client.get('/api/v1/sessions/1/resources').json()[0]['filename'], 'trees.md')
        self.assertEqual(self.client.get(f'/api/v1/sessions/1/resources/{resource_id}').json()['content'],
                         '# Trees\nCourse handout')
        self.assertEqual(self.client.delete(f'/api/v1/sessions/1/resources/{resource_id}').status_code, 403)

        self.as_user(3)
        self.assertEqual(self.client.get('/api/v1/sessions/1/resources').status_code, 403)

    def test_students_cannot_create_or_upload_recordings(self):
        self.as_user(2)
        response = self.client.post('/api/v1/sessions', json={
            'class_course_id': 1, 'title': 'Private recording', 'creator_id': 2,
        })
        self.assertEqual(response.status_code, 403)
        response = self.client.post('/api/v1/sessions/1/chunks/0',
                                    files={'file': ('chunk.pcm', b'audio')})
        self.assertEqual(response.status_code, 403)

    def test_meeting_failure_keeps_pending_then_recovers(self):
        room = self.client.post('/api/v1/meetings', json={'class_course_id': 1, 'name': 'Discussion'}).json()
        self.client.post(f"/api/v1/meetings/{room['id']}/messages", json={'content': 'Explain binary trees'})
        with patch.object(scheduler, 'engine', self.engine), patch.object(scheduler, 'chat', new=AsyncMock(side_effect=RuntimeError('offline'))):
            self.assertEqual(self.client.post(f"/api/v1/meetings/{room['id']}/end").status_code, 502)
        with Session(self.engine) as db:
            saved = db.get(MeetingRoom, room['id'])
            self.assertEqual(saved.processed_message_id, 0)
            self.assertEqual(saved.status, 'ending')
        self.assertEqual(self.client.post(f"/api/v1/meetings/{room['id']}/messages", json={'content': 'late'}).status_code, 409)
        with patch.object(scheduler, 'engine', self.engine), patch.object(scheduler, 'chat', new=AsyncMock(return_value='Summary')):
            response = self.client.post(f"/api/v1/meetings/{room['id']}/end")
            self.assertEqual(response.status_code, 200, response.text)
            self.assertEqual(response.json()['status'], 'ended')
        detail = self.client.get(f"/api/v1/meetings/{room['id']}").json()
        self.assertEqual(len(detail['messages']), 3)
        self.assertEqual(self.client.get(f"/api/v1/meetings/{room['id']}?after={detail['cursor']}").json()['messages'], [])

    def test_end_drains_multiple_batches(self):
        room = self.client.post('/api/v1/meetings', json={'class_course_id': 1, 'name': 'Batch'}).json()
        for i in range(10):
            self.client.post(f"/api/v1/meetings/{room['id']}/messages", json={'content': f'Question {i}'})
        with patch.object(scheduler, 'engine', self.engine), patch.object(scheduler, 'chat', new=AsyncMock(return_value='Summary')):
            response = self.client.post(f"/api/v1/meetings/{room['id']}/end")
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()['status'], 'ended')
        detail = self.client.get(f"/api/v1/meetings/{room['id']}").json()
        self.assertEqual(len(detail['messages']), 14)

    def test_group_chat_is_isolated_from_ai_meeting_scheduler(self):
        ai_room = self.client.post('/api/v1/meetings', json={
            'class_course_id': 1, 'name': 'AI room', 'room_type': 'ai_meeting', 'participant_limit': 6,
        }).json()
        group = self.client.post('/api/v1/meetings', json={
            'class_course_id': 1, 'name': 'Study group', 'room_type': 'group_chat', 'participant_limit': 8,
        }).json()
        self.assertTrue(ai_room['enabled'])
        self.assertFalse(group['enabled'])
        self.assertEqual(
            [row['id'] for row in self.client.get('/api/v1/meetings?class_course_id=1&room_type=ai_meeting').json()],
            [ai_room['id']],
        )
        self.assertEqual(
            [row['id'] for row in self.client.get('/api/v1/meetings?class_course_id=1&room_type=group_chat').json()],
            [group['id']],
        )

        base = f"/api/v1/meetings/{group['id']}"
        with patch.object(scheduler, 'process', new=AsyncMock()) as process:
            response = self.client.post(base + '/messages', json={'content': '@组员A 这只是群聊消息'})
            self.assertEqual(response.status_code, 200, response.text)
            process.assert_not_awaited()
        self.assertEqual(self.client.get(base + '/minutes/export').status_code, 409)
        self.assertEqual(self.client.post(base + '/end').json()['status'], 'ended')
        detail = self.client.get(base).json()
        self.assertEqual([message['role'] for message in detail['messages']], ['user'])

    def test_ai_meeting_direct_mention_uses_named_agent(self):
        room = self.client.post('/api/v1/meetings', json={
            'class_course_id': 1, 'name': 'Direct mention', 'room_type': 'ai_meeting',
        }).json()
        base = f"/api/v1/meetings/{room['id']}"
        with patch.object(scheduler, 'engine', self.engine), \
             patch.object(scheduler, 'chat', new=AsyncMock(return_value='Recorded')):
            response = self.client.post(base + '/messages', json={'content': '@组员C 请提出易错问题'})
        self.assertEqual(response.status_code, 200, response.text)
        messages = self.client.get(base).json()['messages']
        agents = [message for message in messages if message['role'] == 'agent']
        self.assertEqual([message['sender_name'] for message in agents], ['组员C'])

    def test_single_agent_meeting_only_replies_when_mentioned(self):
        room = self.client.post('/api/v1/meetings', json={
            'class_course_id': 1, 'name': 'Single AI', 'agent_mode': 'single',
        }).json()
        self.assertEqual(room['agent_mode'], 'single')
        base = f"/api/v1/meetings/{room['id']}"
        self.client.post(base + '/messages', json={'content': '先记录普通聊天'})
        with patch.object(scheduler, 'engine', self.engine), \
             patch.object(scheduler, 'chat', new=AsyncMock(return_value='Recorded')):
            asyncio.run(scheduler.process(room['id'], force=True))
        self.assertEqual([row['role'] for row in self.client.get(base).json()['messages']], ['user'])

        with patch.object(scheduler, 'engine', self.engine), \
             patch.object(scheduler, 'chat', new=AsyncMock(return_value='AI response')):
            response = self.client.post(base + '/messages', json={'content': '@AI 请解释二叉树'})
        self.assertEqual(response.status_code, 200, response.text)
        agents = [row for row in self.client.get(base).json()['messages'] if row['role'] == 'agent']
        self.assertEqual(agents[-1]['sender_name'], 'AI助手')

    def test_ai_meeting_idle_prompt_is_led_by_agent_a(self):
        room = self.client.post('/api/v1/meetings', json={
            'class_course_id': 1, 'name': 'Idle prompt', 'room_type': 'ai_meeting',
        }).json()
        with Session(self.engine) as db:
            user_message = MeetingMessage(room_id=room['id'], sender_name='Teacher', role='user',
                                          content='我们讨论中序遍历')
            db.add(user_message)
            db.flush()
            db.add(MeetingMessage(room_id=room['id'], sender_name='组员B', role='agent',
                                  content='需要先区分普通二叉树和二叉搜索树'))
            saved = db.get(MeetingRoom, room['id'])
            saved.processed_message_id = user_message.id
            saved.last_activity = datetime.utcnow() - timedelta(seconds=4)
            db.commit()

        reply = AsyncMock(return_value='大家还有哪一步没有想清楚？')
        with patch.object(scheduler, 'engine', self.engine), \
             patch.object(scheduler, 'chat', new=reply), \
             patch.object(scheduler.settings, 'meeting_idle_seconds', 3):
            asyncio.run(scheduler.process(room['id']))

        detail = self.client.get(f"/api/v1/meetings/{room['id']}").json()
        self.assertEqual(detail['messages'][-1]['sender_name'], '组员A')
        self.assertEqual(detail['messages'][-1]['content'], '大家还有哪一步没有想清楚？')
        prompt = reply.await_args.args[1]
        self.assertIn('我们讨论中序遍历', prompt)
        self.assertIn('需要先区分普通二叉树和二叉搜索树', prompt)

    def test_embedding_failure_does_not_leave_evaluation_pending(self):
        self.as_user(2)
        note = self.client.post('/api/v1/notes', json=dict(class_course_id=1, owner_id=2, title='Tree', content='binary trees')).json()
        verdict = json.dumps(dict(relevance=8, correctness=8, structure=8))
        with patch.object(notes, 'chat', new=AsyncMock(return_value=verdict)), patch.object(rag, '_write_vectors', side_effect=HTTPException(502, 'offline')):
            response = self.client.post(f"/api/v1/notes/{note['id']}/share")
        self.assertEqual(response.status_code, 502)
        with Session(self.engine) as db:
            self.assertEqual(db.get(Note, note['id']).quality_status, 'none')
            self.assertEqual(rag.query(1, 'binary trees', db=db), [])

    def test_legacy_publications_and_punctuation_search(self):
        with Session(self.engine) as db:
            db.add(Outline(session_id=1, class_course_id=1, owner_id=-1, markdown='二叉树，遍历。', status='published'))
            db.commit()
            self.assertTrue(rag.query(1, '二叉树', db=db))
            self.assertEqual(rag.query(1, '，。', db=db), [])

    def test_document_pdf_pptx_and_scan(self):
        import fitz
        from io import BytesIO
        from pptx import Presentation
        from pptx.util import Inches
        from app.services import documents
        with fitz.open() as pdf:
            page = pdf.new_page()
            page.insert_text((30, 30), 'Binary tree traversal')
            self.assertIn('Binary tree', documents.extract('lesson.pdf', pdf.tobytes()))
        with fitz.open() as pdf:
            pdf.new_page()
            with patch.object(documents, 'recognize_image', return_value='Scanned tree') as ocr:
                self.assertIn('Scanned tree', documents.extract('scan.pdf', pdf.tobytes()))
                ocr.assert_called_once()
        ppt = Presentation()
        slide = ppt.slides.add_slide(ppt.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = 'Tree slide'
        stream = BytesIO()
        ppt.save(stream)
        self.assertIn('Tree slide', documents.extract('lesson.pptx', stream.getvalue()))
        with self.assertRaises(HTTPException):
            documents.extract('broken.pdf', b'broken')

    def test_background_import_recovery_retry_and_permissions(self):
        import tempfile
        from app.models import DocumentJob, KnowledgeDocument
        from app.services import document_jobs
        with tempfile.TemporaryDirectory() as directory, patch.object(rag.settings, 'data_dir', directory), \
             patch.object(document_jobs, 'engine', self.engine):
            self.as_user(2)
            self.assertEqual(self.client.post('/api/v1/knowledge/1/imports', files={'file': ('test.txt', b'tree')}).status_code, 403)
            self.as_user(1)
            response = self.client.post('/api/v1/knowledge/1/imports', files={'file': ('test.txt', b'binary tree')})
            self.assertEqual(response.status_code, 202)
            key = response.json()['id']
            with Session(self.engine) as db:
                row = db.get(DocumentJob, key)
                row.status = 'processing'  # Crash after claiming the persisted job.
                db.commit()
            document_jobs.run_job(key)
            document_jobs.run_job(key)
            with Session(self.engine) as db:
                self.assertEqual(db.get(DocumentJob, key).status, 'done')
                self.assertEqual(len(db.exec(select(KnowledgeDocument)).all()), 1)
            self.assertFalse(document_jobs.upload_path(key).exists())
            failed = self.client.post('/api/v1/knowledge/1/imports', files={'file': ('broken.pdf', b'invalid')}).json()['id']
            with patch.object(document_jobs.logger, 'exception'):
                document_jobs.run_job(failed)
            self.assertEqual(self.client.get('/api/v1/knowledge/1/jobs').json()[0]['status'], 'failed')
            document_jobs.upload_path(failed).write_bytes(b'fixed content')
            with Session(self.engine) as db:
                row = db.get(DocumentJob, failed)
                row.filename = 'fixed.txt'
                db.commit()
            queued = self.client.post('/api/v1/knowledge/1/imports', files={'file': ('waiting.txt', b'waiting')}).json()['id']
            with patch.object(rag.settings, 'document_queue_limit', 1):
                self.assertEqual(self.client.post(f'/api/v1/knowledge/1/jobs/{failed}/retry').status_code, 429)
            document_jobs.run_job(queued)
            self.assertEqual(self.client.post(f'/api/v1/knowledge/1/jobs/{failed}/retry').status_code, 200)
            with patch('pathlib.Path.unlink', side_effect=PermissionError('busy')), \
                 patch.object(document_jobs.logger, 'warning') as warning:
                document_jobs.run_job(failed)
                warning.assert_called_once()
            with Session(self.engine) as db:
                self.assertEqual(db.get(DocumentJob, failed).status, 'done')
                self.assertIsNotNone(db.get(KnowledgeDocument, f'1:import_{failed}'))
            self.as_user(3)
            self.assertEqual(self.client.get('/api/v1/knowledge/1/jobs').status_code, 403)

    def test_minutes_export_requires_participation(self):
        room = self.client.post('/api/v1/meetings', json={'class_course_id': 1, 'name': 'Minutes'}).json()['id']
        path = f'/api/v1/meetings/{room}/minutes/export'
        self.assertEqual(self.client.get(path).status_code, 409)
        with Session(self.engine) as db:
            row = db.get(MeetingRoom, room)
            row.minutes = '## Conclusion\nBinary trees'
            db.commit()
        response = self.client.get(path)
        self.assertIn('Binary trees', response.text)
        self.assertIn('attachment;', response.headers['content-disposition'])
        self.as_user(2)
        self.assertEqual(self.client.get(path).status_code, 403)
        self.client.post(f'/api/v1/meetings/{room}/join')
        self.assertEqual(self.client.get(path).status_code, 200)

    def test_many_meetings_concurrently_commit_once(self):
        with Session(self.engine) as db:
            for room_id in range(100, 112):
                db.add(MeetingRoom(id=room_id, class_course_id=1, creator_id=1, name=f'Room {room_id}'))
                db.add(MeetingMessage(room_id=room_id, sender_id=1, sender_name='Teacher', content='@分析师 tree'))
            db.commit()
        active, peak = 0, 0
        async def reply(*args):
            nonlocal active, peak
            active += 1
            peak = max(peak, active)
            await asyncio.sleep(0.005)
            active -= 1
            return '[]' if 'JSON' in args[0] else 'Result'
        async def run():
            gate = asyncio.Semaphore(4)
            async def one(room_id):
                async with gate: await scheduler.process(room_id, force=True)
            await asyncio.gather(*(one(room_id) for room_id in range(100, 112)))
            await asyncio.gather(*(scheduler.process(100, force=True) for _ in range(5)))
        with patch.object(scheduler, 'engine', self.engine), patch.object(scheduler, 'chat', reply), \
             patch.object(rag, 'query', return_value=[]):
            asyncio.run(run())
        self.assertGreater(peak, 1)
        self.assertLessEqual(peak, 4)
        with Session(self.engine) as db:
            for room_id in range(100, 112):
                self.assertTrue(db.get(MeetingRoom, room_id).minutes.strip())
                replies = db.exec(select(MeetingMessage).where(MeetingMessage.room_id == room_id, MeetingMessage.role == 'agent')).all()
                self.assertLessEqual(len(replies), 2)
                self.assertGreater(len(replies), 0)

    def test_document_group_chart_and_real_local_ocr(self):
        from io import BytesIO
        from PIL import Image, ImageDraw, ImageFont
        from pptx import Presentation
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE
        from pptx.util import Inches
        from app.services import documents
        image = Image.new('RGB', (800, 180), 'white')
        draw = ImageDraw.Draw(image)
        draw.text((20, 30), 'Binary tree 2026', fill='black', font=ImageFont.truetype('DejaVuSans.ttf', 40) if __import__('os').name != 'nt' else ImageFont.truetype('arial.ttf', 40))
        stream = BytesIO()
        image.save(stream, format='PNG')
        with patch.object(rag.settings, 'document_ocr_provider', 'local'):
            self.assertIn('2026', documents.recognize_image(stream.getvalue()))
        ppt = Presentation()
        slide = ppt.slides.add_slide(ppt.slide_layouts[6])
        group = slide.shapes.add_group_shape()
        group.shapes.add_textbox(0, 0, Inches(3), Inches(1)).text = 'Nested group text'
        chart = CategoryChartData()
        chart.categories = ['A', 'B']
        chart.add_series('Scores', [12, 34])
        slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, 0, Inches(2), Inches(4), Inches(3), chart)
        slide.notes_slide.notes_text_frame.text = 'Speaker note'
        stream = BytesIO()
        ppt.save(stream)
        result = documents.extract('chart.pptx', stream.getvalue())
        for expected in ('Nested group text', 'Scores', 'A: 12', 'B: 34', 'Speaker note'):
            self.assertIn(expected, result)

    def test_chat_auth_stream_metadata_and_deleted_session(self):
        from app.models import ChatSession, ChatMessage
        response = self.client.post('/api/v1/chat/sessions', json={'user_id': 1, 'class_course_id': 1})
        chat_id = response.json()['id']
        self.as_user(2)
        self.assertEqual(self.client.get(f'/api/v1/chat/sessions/{chat_id}/messages').status_code, 403)
        self.assertEqual(self.client.post('/api/v1/chat/ask', json={'chat_session_id': chat_id, 'question': 'tree'}).status_code, 403)
        self.as_user(1)
        async def reply(*args):
            yield 'Answer'
        with patch.object(chat, 'engine', self.engine), patch.object(chat, 'chat_stream', reply):
            response = self.client.post('/api/v1/chat/ask', json={'chat_session_id': chat_id, 'question': 'tree'})
        self.assertIn('event: done', response.text)
        self.assertIn('retrieval_mode', response.text)
        async def delete_during_stream(*args):
            with Session(self.engine) as db:
                db.delete(db.get(ChatSession, chat_id))
                db.commit()
            yield 'Orphan answer'
        with patch.object(chat, 'engine', self.engine), patch.object(chat, 'chat_stream', delete_during_stream):
            self.client.post('/api/v1/chat/ask', json={'chat_session_id': chat_id, 'question': 'tree'})
        with Session(self.engine) as db:
            self.assertIsNone(db.exec(select(ChatMessage).where(ChatMessage.content == 'Orphan answer')).first())

    def test_personal_copy_survives_official_withdrawal(self):
        with Session(self.engine) as db:
            db.add(Outline(session_id=1, class_course_id=1, owner_id=-1, markdown='Official', status='published'))
            db.commit()
        self.as_user(2)
        self.client.put('/api/v1/sessions/1/outline?user_id=2', json={'markdown': 'My saved version'})
        self.as_user(1)
        self.client.post('/api/v1/sessions/1/outline/review', json={'user_id': 1, 'action': 'reject'})
        self.as_user(2)
        self.assertEqual(self.client.get('/api/v1/sessions/1/outline').json()['markdown'], 'My saved version')
        self.assertEqual(self.client.get('/api/v1/sessions?class_course_id=1').json()[0]['outline_kind'], 'personal')

    def test_vector_search_filters_deleted_versions(self):
        import tempfile
        # Deterministic embeddings exercise the actual Chroma query contract without an external API.
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as directory:
            with patch.object(rag.settings, 'data_dir', directory), patch.object(rag.settings, 'embed_api_key', 'test'), \
                 patch.object(rag.settings, 'embed_model', 'test-model'), patch.object(rag, 'embed', side_effect=lambda texts: [[1.0, 0.0, 0.0] for _ in texts]):
                with Session(self.engine) as db:
                    rag.add_document(1, 'manual-test', 'binary tree', 'old', db=db)
                    db.commit()
                    hits = rag.query(1, 'tree', db=db)
                    self.assertEqual(hits[0]['retrieval_mode'], 'hybrid')
                    rag.add_document(1, 'manual-test', 'updated tree', 'new', db=db)
                    db.commit()
                    self.assertEqual({hit['source'] for hit in rag.query(1, 'tree', db=db)}, {'new'})
                    self.assertEqual(rag.cleanup(db), 1)
                    self.assertEqual(rag._collection(1).count(), 1)
                    self.assertEqual({hit['source'] for hit in rag.query(1, 'tree', db=db)}, {'new'})
                    self.assertEqual(rag.query(2, 'tree', db=db), [])
                    rag.delete_document(1, 'manual-test', db=db)
                    db.commit()
                    self.assertEqual(rag.query(1, 'tree', db=db), [])
                    self.assertEqual(rag.cleanup(db), 1)
                    self.assertEqual(rag._collection(1).count(), 0)

    def account_identity(self):
        identity = dict(account_type='school', school_name='University', student_number='2026001')
        with Session(self.engine) as db:
            db.add(Account(user_id=1, login_key=auth.identity(auth.Identity(**identity)), account_type='school',
                           password_hash=hash_password('old-password'), email='teacher@example.com'))
            db.commit()
        return identity

    def test_password_reset_consumes_code_and_revokes_all_sessions(self):
        identity = self.account_identity()
        with patch.object(auth, 'deliver_account_code') as delivery:
            response = self.client.post('/api/v1/auth/password-reset/code', json=identity)
            self.assertEqual(response.status_code, 200, response.text)
            code = delivery.call_args.args[1]
        with Session(self.engine) as db:
            saved = db.get(AccountCode, 'reset:1')
            self.assertNotEqual(saved.code_hash, code)
        reset = {**identity, 'code': code, 'new_password': 'new-password'}
        response = self.client.post('/api/v1/auth/password-reset', json=reset)
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(self.client.get('/api/v1/auth/me').status_code, 401)
        self.assertEqual(self.client.post('/api/v1/auth/password-reset', json=reset).status_code, 400)
        self.assertEqual(self.client.post('/api/v1/auth/login', json={**identity, 'password': 'old-password'}).status_code, 401)
        self.assertEqual(self.client.post('/api/v1/auth/login', json={**identity, 'password': 'new-password'}).status_code, 200)

    def test_reset_code_expiry_attempt_limit_and_identity_isolation(self):
        identity = self.account_identity()
        with patch.object(auth, 'deliver_account_code') as delivery:
            self.client.post('/api/v1/auth/password-reset/code', json=identity)
            code = delivery.call_args.args[1]
        reset = {**identity, 'code': code, 'new_password': 'new-password'}
        self.assertEqual(self.client.post('/api/v1/auth/password-reset', json={**reset, 'student_number': 'other'}).status_code, 400)
        wrong = '000000' if code != '000000' else '111111'
        for _ in range(5):
            self.assertEqual(self.client.post('/api/v1/auth/password-reset', json={**reset, 'code': wrong}).status_code, 400)
        self.assertEqual(self.client.post('/api/v1/auth/password-reset', json=reset).status_code, 400)
        with Session(self.engine) as db:
            row = db.get(AccountCode, 'reset:1')
            row.attempts = 0
            row.expires_at = datetime.utcnow() - timedelta(seconds=1)
            db.commit()
        self.assertEqual(self.client.post('/api/v1/auth/password-reset', json=reset).status_code, 400)

    def test_password_change_requires_old_password_and_invalidates_reset(self):
        identity = self.account_identity()
        with patch.object(auth, 'deliver_account_code') as delivery:
            self.client.post('/api/v1/auth/password-reset/code', json=identity)
            code = delivery.call_args.args[1]
        self.assertEqual(self.client.post('/api/v1/auth/password', json={'current_password': 'incorrect', 'new_password': 'new-password'}).status_code, 400)
        self.assertEqual(self.client.post('/api/v1/auth/password', json={'current_password': 'old-password', 'new_password': 'new-password'}).status_code, 200)
        self.assertEqual(self.client.get('/api/v1/auth/me').status_code, 401)
        self.assertEqual(self.client.post('/api/v1/auth/password-reset', json={**identity, 'code': code, 'new_password': 'third-password'}).status_code, 400)

    def test_recovery_email_binding_is_verified_and_codes_are_purpose_bound(self):
        identity = self.account_identity()
        body = {'email': 'new@example.com', 'current_password': 'old-password'}
        with patch.object(auth, 'deliver_account_code') as delivery:
            self.assertEqual(self.client.post('/api/v1/auth/recovery-email/code', json=body).status_code, 200)
            code = delivery.call_args.args[1]
        self.assertEqual(self.client.post('/api/v1/auth/password-reset', json={**identity, 'code': code, 'new_password': 'new-password'}).status_code, 400)
        self.assertEqual(self.client.post('/api/v1/auth/recovery-email', json={**body, 'email': 'other@example.com', 'code': code}).status_code, 400)
        response = self.client.post('/api/v1/auth/recovery-email', json={**body, 'code': code})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()['email'], 'new@example.com')
        self.assertEqual(self.client.post('/api/v1/auth/recovery-email', json={**body, 'code': code}).status_code, 400)

    def test_unknown_reset_account_does_not_send_email(self):
        with patch.object(auth, 'deliver_account_code') as delivery:
            response = self.client.post('/api/v1/auth/password-reset/code', json={'email': 'unknown@example.com'})
            self.assertEqual(response.status_code, 200)
            delivery.assert_not_called()

    def test_meeting_members_permissions_removal_reinvite_and_leave(self):
        room = self.client.post('/api/v1/meetings', json={'class_course_id': 1, 'name': 'Members', 'participant_limit': 2}).json()
        base = f"/api/v1/meetings/{room['id']}"
        self.assertEqual(self.client.post(base + '/members', json={'user_id': 3}).status_code, 403)
        self.assertEqual(self.client.post(base + '/members', json={'user_id': 2}).status_code, 200)
        self.as_user(2)
        self.assertEqual(self.client.get(base + '/members').status_code, 403)
        self.assertEqual(self.client.delete(base + '/members/1').status_code, 403)
        self.as_user(1)
        self.assertEqual(self.client.delete(base + '/members/1').status_code, 409)
        self.assertEqual(self.client.delete(base + '/members/2').status_code, 200)
        self.as_user(2)
        self.assertEqual(self.client.post(base + '/join').status_code, 403)
        self.assertEqual(self.client.get(base).status_code, 403)
        self.assertEqual(self.client.post(base + '/messages', json={'content': 'blocked'}).status_code, 403)
        self.as_user(1)
        self.assertEqual(self.client.post(base + '/members', json={'user_id': 2}).status_code, 200)
        self.assertEqual(self.client.post(base + '/leave').status_code, 409)
        self.assertEqual(self.client.post(base + '/owner', json={'user_id': 2}).status_code, 200)
        self.assertEqual(self.client.post(base + '/leave').status_code, 200)
        self.as_user(2)
        self.assertEqual(self.client.get(base + '/members').status_code, 200)

    def test_ended_meeting_does_not_admit_new_members(self):
        room = self.client.post('/api/v1/meetings', json={'class_course_id': 1, 'name': 'Closed'}).json()
        base = f"/api/v1/meetings/{room['id']}"
        with patch.object(scheduler, 'engine', self.engine):
            self.assertEqual(self.client.post(base + '/end').status_code, 200)
        self.assertEqual(self.client.post(base + '/members', json={'user_id': 2}).status_code, 409)
        self.as_user(2)
        self.assertEqual(self.client.post(base + '/join').status_code, 409)


if __name__ == '__main__':
    unittest.main()
